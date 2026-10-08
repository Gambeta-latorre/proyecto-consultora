"""Vista previa aproximada de un .pptx (sin PowerPoint): dibuja formas, imágenes y texto con Pillow."""
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image, ImageDraw, ImageFont

FD = "C:/Windows/Fonts/"
FF = {("calibri", 0): "calibri.ttf", ("calibri", 1): "calibrib.ttf", ("arial", 0): "arial.ttf", ("arial", 1): "arialbd.ttf",
      ("cambria", 0): "cambria.ttc", ("cambria", 1): "cambriab.ttf", ("segoe ui", 0): "segoeui.ttf", ("segoe ui", 1): "segoeuib.ttf"}
K = 100  # px por pulgada
cache = {}


def font(name, bold, pt):
    n = (name or "calibri").lower()
    f = FF.get((n, 1 if bold else 0), FF[("calibri", 1 if bold else 0)])
    key = (f, round(pt * K / 72))
    if key not in cache:
        cache[key] = ImageFont.truetype(FD + f, max(6, round(pt * K / 72)))
    return cache[key]


def col(c, default=(0, 0, 0)):
    try:
        return tuple(c.rgb)
    except Exception:
        return default


def draw_text(d, shp, x, y, w, h):
    tf = shp.text_frame
    ml = (tf.margin_left or 91440) / 914400 * K
    mr = (tf.margin_right or 91440) / 914400 * K
    mt = (tf.margin_top or 45720) / 914400 * K
    lines = []
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        r0 = runs[0]
        sz = r0.font.size.pt if r0.font.size else 18
        f = font(r0.font.name, r0.font.bold, sz)
        try:
            c = tuple(r0.font.color.rgb)
        except Exception:
            c = (0, 0, 0)
        txt = "".join(r.text for r in runs)
        pPr = p._p.pPr
        bullet = pPr is not None and pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None
        avail = w - ml - mr - (28 if bullet else 0)
        words, cur, wrapped = txt.split(" "), "", []
        for wd in words:
            t = (cur + " " + wd).strip()
            if d.textlength(t, font=f) <= avail or not cur:
                cur = t
            else:
                wrapped.append(cur)
                cur = wd
        wrapped.append(cur)
        ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        sa = (p.space_after.pt if p.space_after is not None else 0) * K / 72
        lines.append((wrapped, f, c, sz * K / 72 * 1.2 * ls, sa, p.alignment, bullet))
    total = sum(len(l[0]) * l[3] + l[4] for l in lines)
    anchor = tf.vertical_anchor
    ay = y + mt
    if anchor is not None and int(anchor) == 3:
        ay = y + (h - total) / 2
    elif anchor is not None and int(anchor) == 4:
        ay = y + h - total
    for wrapped, f, c, lh, sa, al, bullet in lines:
        for i, ln in enumerate(wrapped):
            tw = d.textlength(ln, font=f)
            if al is not None and int(al) == 2:
                tx = x + (w - tw) / 2
            elif al is not None and int(al) == 3:
                tx = x + w - mr - tw
            else:
                tx = x + ml + (28 if bullet else 0)
            if bullet and i == 0:
                d.text((x + ml + 6, ay), "•", font=f, fill=c)
            d.text((tx, ay), ln, font=f, fill=c)
            ay += lh
        ay += sa


def render(path, out_prefix, only=None):
    prs = Presentation(path)
    W, H = prs.slide_width / 914400 * K, prs.slide_height / 914400 * K
    outs = []
    for i, sl in enumerate(prs.slides, 1):
        if only and i not in only:
            continue
        im = Image.new("RGB", (int(W), int(H)), (255, 255, 255))
        d = ImageDraw.Draw(im, "RGBA")
        try:
            if sl.background.fill.type == 1:
                d.rectangle([0, 0, W, H], fill=col(sl.background.fill.fore_color, (255, 255, 255)))
        except Exception:
            pass
        for shp in sl.shapes:
            x, y, w, h = [v / 914400 * K for v in (shp.left, shp.top, shp.width, shp.height)]
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pic = Image.open(shp.image.blob if False else __import__("io").BytesIO(shp.image.blob)).convert("RGBA")
                pic = pic.resize((max(1, int(w)), max(1, int(h))))
                im.paste(pic, (int(x), int(y)), pic)
                continue
            if getattr(shp, "has_table", False) and shp.has_table:
                tb = shp.table
                cy = y
                for r in tb.rows:
                    rh = r.height / 914400 * K
                    cx = x
                    for ci, c in enumerate(r.cells):
                        cw = tb.columns[ci].width / 914400 * K
                        fillc = None
                        try:
                            if c.fill.type == 1:
                                fillc = tuple(c.fill.fore_color.rgb)
                        except Exception:
                            pass
                        d.rectangle([cx, cy, cx + cw, cy + rh], fill=fillc, outline=(200, 200, 200))
                        class _S:  # adaptar celda a draw_text
                            text_frame = c.text_frame
                            has_text_frame = True
                        c.text_frame.margin_left = c.margin_left
                        c.text_frame.margin_right = c.margin_right
                        c.text_frame.margin_top = c.margin_top
                        draw_text(d, _S, cx, cy, cw, rh)
                        cx += cw
                    cy += rh
                continue
            if shp.has_chart if hasattr(shp, "has_chart") else False:
                d.rectangle([x, y, x + w, y + h], outline=(150, 150, 150))
                d.text((x + 10, y + 10), "[gráfico nativo]", fill=(100, 100, 100), font=font("calibri", 0, 12))
                continue
            if shp.shape_type == MSO_SHAPE_TYPE.FREEFORM:
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                path = shp._element.spPr.find(".//a:path", ns)
                pw, ph = int(path.get("w")) or 1, int(path.get("h")) or 1
                pts = [(x + int(pt.get("x")) / pw * w, y + int(pt.get("y")) / ph * h) for pt in path.findall(".//a:pt", ns)]
                d.line(pts, fill=(40, 40, 40), width=3)
                (x0_, y0_), (x1_, y1_) = pts[-2], pts[-1]
                import math
                ang = math.atan2(y1_ - y0_, x1_ - x0_)
                d.polygon([(x1_, y1_), (x1_ - 14 * math.cos(ang - .4), y1_ - 14 * math.sin(ang - .4)), (x1_ - 14 * math.cos(ang + .4), y1_ - 14 * math.sin(ang + .4))], fill=(40, 40, 40))
                continue
            if shp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill = None
                try:
                    if shp.fill.type == 1:
                        fill = col(shp.fill.fore_color)
                except Exception:
                    pass
                line = None
                try:
                    if shp.line.fill.type == 1:
                        line = col(shp.line.color)
                except Exception:
                    pass
                kind = shp.auto_shape_type
                box = [x, y, x + w, y + h]
                if "OVAL" in str(kind):
                    d.ellipse(box, fill=fill, outline=line, width=2)
                elif "ROUNDED" in str(kind):
                    d.rounded_rectangle(box, radius=min(w, h) * (shp.adjustments[0] if len(shp.adjustments) else .1), fill=fill, outline=line, width=2)
                else:
                    d.rectangle(box, fill=fill, outline=line, width=2)
            if shp.has_text_frame and shp.text_frame.text.strip():
                draw_text(d, shp, x, y, w, h)
        o = f"{out_prefix}{i:02d}.png"
        im.save(o)
        outs.append(o)
    return outs


if __name__ == "__main__":
    print("\n".join(render(sys.argv[1], sys.argv[2])))


def montage(files, out, cols=2, scale=0.7):
    ims = [Image.open(f) for f in files]
    w, h = int(ims[0].width * scale), int(ims[0].height * scale)
    rows = (len(ims) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * w + (cols + 1) * 8, rows * h + (rows + 1) * 8), (120, 120, 120))
    for i, im in enumerate(ims):
        r, c = divmod(i, cols)
        sheet.paste(im.resize((w, h)), (8 + c * (w + 8), 8 + r * (h + 8)))
    sheet.save(out)
    return out

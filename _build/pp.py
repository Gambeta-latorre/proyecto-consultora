"""Helpers para armar presentaciones con python-pptx y verificar que el texto entre en su caja."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import ImageFont
import copy

FD = "C:/Windows/Fonts/"
FONTFILES = {
    ("Calibri", False): "calibri.ttf", ("Calibri", True): "calibrib.ttf",
    ("Cambria", False): "cambria.ttc", ("Cambria", True): "cambriab.ttf",
    ("Arial", False): "arial.ttf", ("Arial", True): "arialbd.ttf",
}
_fc = {}
WARN = []


def rgb(h):
    return RGBColor.from_string(h)


def _font(name, bold, size):
    key = (name, bold, size)
    if key not in _fc:
        _fc[key] = ImageFont.truetype(FD + FONTFILES[(name, bold)], int(size * 10))
    return _fc[key]


def text_height(paras, width_in, default_size, font, bold=False, spacing=1.2, para_gap=0):
    """Altura estimada (pulgadas) de una lista de párrafos [(texto, size, bold)] en un ancho dado."""
    total = 0
    for p in paras:
        txt, size, b = p
        f = _font(font, b or bold, size)
        maxw = width_in * 72 * 10
        lines = 0
        for seg in str(txt).split("\n"):
            cur = ""
            n = 1
            for w in seg.split(" "):
                t = (cur + " " + w).strip()
                if f.getlength(t) <= maxw or not cur:
                    cur = t
                else:
                    n += 1
                    cur = w
            lines += n
        total += lines * size * spacing + para_gap
    return total / 72


class Deck:
    def __init__(self, T):
        self.T = T
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.W, self.H = 13.333, 7.5
        self.n = 0

    def slide(self, dark=False, title=None, title_color=None, size=34, notes=None, tx=0.7, ty=0.45, tw=11.9, th=0.9):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        self.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(self.T["dark"] if dark else self.T["light"])
        t = s.shapes.title
        if title is None:
            t._element.getparent().remove(t._element)
        else:
            t.left, t.top, t.width, t.height = Inches(tx), Inches(ty), Inches(tw), Inches(th)
            tf = t.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = 0
            p = tf.paragraphs[0]
            p.text = title
            p.alignment = PP_ALIGN.LEFT
            r = p.runs[0]
            r.font.name = self.T["head"]
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = rgb(title_color or (self.T["on_dark"] if dark else self.T["ink"]))
            h = text_height([(title, size, True)], tw, size, self.T["head"])
            if h > th + 0.02:
                WARN.append(f"slide {self.n}: título desborda ({h:.2f} > {th})")
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        s._dark = dark
        return s

    def text(self, s, x, y, w, h, content, size=16, bold=False, color=None, align="l", anchor="t", font=None,
             spacing=1.15, para_gap=4, bullets=False, check=True, rtl=False, italic=False):
        """content: str | list[str | (str, dict)]. dict admite size, bold, color, italic."""
        font = font or self.T["body"]
        color = color or (self.T["on_dark"] if getattr(s, "_dark", False) else self.T["ink"])
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        items = content if isinstance(content, list) else [content]
        meas = []
        for i, it in enumerate(items):
            txt, o = (it, {}) if isinstance(it, str) else it
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[o.get("align", align)]
            p.line_spacing = spacing
            p.space_after = Pt(o.get("gap", para_gap))
            r = p.add_run()
            r.text = txt
            sz = o.get("size", size)
            r.font.size = Pt(sz)
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", italic)
            r.font.name = o.get("font", font)
            r.font.color.rgb = rgb(o.get("color", color))
            if rtl or o.get("rtl"):
                p._p.get_or_add_pPr().set("rtl", "1")
                r._r.get_or_add_rPr().set("lang", "ar-SA")
                rPr = r._r.get_or_add_rPr()
                cs = rPr.makeelement(qn("a:cs"), {"typeface": o.get("font", font)})
                rPr.append(cs)
            if bullets and not o.get("nobullet"):
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(int(Inches(0.28))))
                pPr.set("indent", str(int(-Inches(0.22))))
                bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
                pPr.append(bu)
            meas.append((txt, sz, o.get("bold", bold)))
        if check:
            wi = w - 0.08 - (0.28 if bullets else 0)
            need = text_height(meas, wi, size, font, spacing=spacing * 1.0, para_gap=para_gap) + 0.06
            if need > h + 0.03:
                WARN.append(f"slide {self.n}: texto desborda ({need:.2f} > {h}) -> {str(items[0])[:50]}")
        return tb

    def rect(self, s, x, y, w, h, fill=None, line=None, shape="round", lw=1.5, radius=0.08, shadow=False):
        shp = s.shapes.add_shape({"round": MSO_SHAPE.ROUNDED_RECTANGLE, "rect": MSO_SHAPE.RECTANGLE, "oval": MSO_SHAPE.OVAL}[shape],
                                 Inches(x), Inches(y), Inches(w), Inches(h))
        if shape == "round":
            shp.adjustments[0] = min(0.5, radius / min(w, h))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(lw)
        else:
            shp.line.fill.background()
        if not shadow:
            sp = shp._element.spPr
            ef = sp.makeelement(qn("a:effectLst"), {})
            sp.append(ef)
        shp.text_frame.text = ""
        return shp

    def num(self, s, x, y, d, n, fill, color="FFFFFF", size=16):
        c = self.rect(s, x, y, d, d, fill=fill, shape="oval")
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(n)
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.name = self.T["body"]
        r.font.color.rgb = rgb(color)
        return c

    def image(self, s, path, x, y, w=None, h=None, alt=None):
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), **kw)
        if alt:
            pic._element.nvPicPr.cNvPr.set("descr", alt)
        return pic

    def save(self, path):
        self.prs.save(path)
        return path

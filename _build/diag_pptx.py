"""Genera una presentación con los diagramas como FORMAS EDITABLES (cajas, rombos, flechas), lista para importar a Google Slides."""
import sys
sys.path.insert(0, "_build")
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
import diag, diag_complex
from diag import Diagram
from diagrams_meta import DIAGS

SW, SH = 20.0, 11.25            # lienzo 16:9 grande: así los textos quedan en 10 a 14 pt
captured = {}


def fake_save(self, path):
    captured[path.split("/")[-1]] = self
    return path


Diagram.save = fake_save
diag.build_simple()
for f in (diag_complex.c1, diag_complex.c2, diag_complex.c3, diag_complex.c4, diag_complex.c5, diag_complex.c6):
    f()


def rgb(t):
    return RGBColor(*t)


def hexs(t):
    return "%02X%02X%02X" % t


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
BLANK_TITLE = prs.slide_layouts[5]


def set_text(shape, text, size, color, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font="Calibri"):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = rgb(color) if isinstance(color, tuple) else RGBColor.from_string(color)


def add_box(slide, kind, x, y, w, h, fill=None, line=None, lw=2.0, name=None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
    sp = shp._element.spPr
    sp.append(sp.makeelement(qn("a:effectLst"), {}))     # sin sombra
    if name:
        shp.name = name
    return shp


def polyline(slide, pts, color, width_pt, dashed=False, name=None):
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=False)
    shp = fb.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = rgb(color)
    shp.line.width = Pt(width_pt)
    if dashed:
        shp.line.dash_style = MSO_LINE.DASH
    ln = shp.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    sp = shp._element.spPr
    sp.append(sp.makeelement(qn("a:effectLst"), {}))
    if name:
        shp.name = name
    return shp


def edge_points(D, e):
    a, pa, b, pb, label, via, lpos, dash = e
    p1, p2 = D.port(a, pa), D.port(b, pb)
    if via:
        return [p1] + via + [p2]
    hz = lambda p: p in "lr"
    if hz(pa) and hz(pb):
        if abs(p1[1] - p2[1]) < 2:
            return [p1, p2]
        mx = (p1[0] + p2[0]) / 2
        return [p1, (mx, p1[1]), (mx, p2[1]), p2]
    if not hz(pa) and not hz(pb):
        if abs(p1[0] - p2[0]) < 2:
            return [p1, p2]
        my = (p1[1] + p2[1]) / 2
        return [p1, (p1[0], my), (p2[0], my), p2]
    if hz(pa):
        return [p1, (p2[0], p1[1]), p2]
    return [p1, (p1[0], p2[1]), p2]


def draw_diagram(slide, D, area):
    ax, ay, aw, ah = area
    top = 100                                           # franja de título del PNG (se omite: el título va en la diapositiva)
    k = min(aw / D.W, ah / (D.H - top))
    ox = ax + (aw - D.W * k) / 2
    X = lambda x: ox + x * k
    Y = lambda y: ay + (y - top) * k
    t = D.t
    pt = lambda px: px * k * 72
    # carriles
    if D.lanes:
        for i, name in enumerate(D.lanes):
            y0 = D.top + i * D.lane_h
            add_box(slide, MSO_SHAPE.RECTANGLE, X(0), Y(y0), D.W * k, D.lane_h * k, fill=t["tint"] if i % 2 == 0 else t["tint2"], name=f"Carril {i+1}")
            hdr = add_box(slide, MSO_SHAPE.RECTANGLE, X(0), Y(y0), (D.hdr - 20) * k, D.lane_h * k, fill=t["pri"] if i % 2 == 0 else t["lane2"], name=f"Título carril {i+1}")
            set_text(hdr, name, pt(26), t["on"], bold=True)
    # flechas (debajo de las cajas)
    for n, e in enumerate(D.edges, 1):
        pts = [(X(x), Y(y)) for x, y in edge_points(D, e)]
        polyline(slide, pts, t["line"], max(1.5, 5 * k * 72 * 0.6), dashed=e[7], name=f"Flecha {n}")
    # cajas
    for id_, (cx, cy, w, h, kind, text, num) in D.nodes.items():
        x, y, ww, hh = X(cx - w / 2), Y(cy - h / 2), w * k, h * k
        if kind == "proc":
            shp = add_box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, ww, hh, fill=(255, 255, 255), line=t["pri"], lw=max(1.5, 4 * k * 72 * 0.6), name=f"Paso {id_}")
            shp.adjustments[0] = 0.12
            set_text(shp, text, pt(D.fs), t["txt"])
        elif kind == "term":
            shp = add_box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, ww, hh, fill=t["pri"], name=f"Inicio/fin {id_}")
            shp.adjustments[0] = 0.5
            set_text(shp, text, pt(D.fs), t["on"], bold=True)
        else:
            shp = add_box(slide, MSO_SHAPE.DIAMOND, x, y, ww, hh, fill=t["dec"], line=t["acc"], lw=max(1.5, 5 * k * 72 * 0.6), name=f"Decisión {id_}")
            set_text(shp, text, pt(D.fs) * 0.95, t["txt"], bold=True)
            shp.text_frame.margin_left = shp.text_frame.margin_right = Inches(w * k * 0.2)
        if num is not None:
            r = 26 * k
            b = add_box(slide, MSO_SHAPE.OVAL, X(cx - w / 2) - r + 4 * k, Y(cy - h / 2) - r + 4 * k, 2 * r, 2 * r, fill=t["acc"], name=f"Número {num}")
            set_text(b, str(num), pt(28), (255, 255, 255) if t["acc"] == (208, 2, 27) else t["txt"], bold=True)
    # etiquetas SÍ / NO
    for e in D.edges:
        label = e[4]
        if not label:
            continue
        pts = edge_points(D, e)
        (sx, sy), (ex, ey) = pts[0], pts[1]
        lpos = e[6]
        if abs(ex - sx) < abs(ey - sy):
            lx, ly = sx + 34, sy + (ey - sy) * 0.5 + lpos
        else:
            lx, ly = sx + (ex - sx) * 0.5 + lpos, sy - 26
        w = (len(label) * 15 + 30) * k
        h = 38 * k
        fill = t["acc"] if t["acc"] == (208, 2, 27) else t["pri"]
        b = add_box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, X(lx) - w / 2, Y(ly) - h / 2, w, h, fill=fill, name=f"Etiqueta {label}")
        b.adjustments[0] = 0.3
        set_text(b, label, pt(24), (255, 255, 255), bold=True)


def add_title_slide(title, subtitle=None, dark=True, bg="173A2B"):
    s = prs.slides.add_slide(BLANK_TITLE)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(1.2), Inches(3.6), Inches(17.6), Inches(1.8)
    t.text_frame.text = title
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.runs[0]
    r.font.size, r.font.bold, r.font.name = Pt(60), True, "Cambria"
    r.font.color.rgb = RGBColor.from_string("C9A227")
    if subtitle:
        tb = s.shapes.add_textbox(Inches(1.2), Inches(5.6), Inches(17.6), Inches(2))
        tb.text_frame.word_wrap = True
        tb.text_frame.text = subtitle
        for r in tb.text_frame.paragraphs[0].runs:
            r.font.size, r.font.name = Pt(26), "Calibri"
            r.font.color.rgb = RGBColor.from_string("DCE9E1")
    return s


add_title_slide("Diagramas de flujo ASHAB y Rumbo Global",
                "Editables: cada caja, rombo y flecha es un objeto que podés mover, cambiar de color o borrar. Archivo > Descargar para exportar a PowerPoint o PDF.")
for tipo, fname, titulo, cap in DIAGS:
    D = captured.get(fname)
    s = prs.slides.add_slide(BLANK_TITLE)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.35), Inches(18.8), Inches(1.0)
    t.text_frame.text = titulo
    tp = t.text_frame.paragraphs[0]
    tp.alignment = PP_ALIGN.LEFT
    tr = tp.runs[0]
    tr.font.size, tr.font.bold, tr.font.name = Pt(36), True, "Cambria"
    tr.font.color.rgb = RGBColor.from_string("111111" if D.t["pri"] == (17, 17, 17) else "1F4D3A")
    draw_diagram(s, D, (0.5, 1.5, 19.0, 7.9))
    tb = s.shapes.add_textbox(Inches(0.6), Inches(9.6), Inches(18.8), Inches(1.4))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = cap
    for r in tb.text_frame.paragraphs[0].runs:
        r.font.size, r.font.name = Pt(18), "Calibri"
        r.font.color.rgb = RGBColor.from_string("566259")
    s.notes_slide.notes_text_frame.text = cap

out = "02_Presentaciones/04_Diagramas_Editables_para_Google_Slides.pptx"
prs.save(out)
print(out, len(prs.slides), "diapositivas")

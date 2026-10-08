"""Helpers para generar documentos Word con python-docx."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

THEMES = {
    "rg": dict(pri="111111", acc="D0021B", soft="F3F3F3", head="Arial"),
    "ash": dict(pri="1F4D3A", acc="C9A227", soft="F3F7F4", head="Cambria"),
}


def rgb(h):
    return RGBColor.from_string(h)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def set_cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def borders(table, color="C8D3CB"):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


class Doc:
    def __init__(self, theme, header_text=""):
        self.t = THEMES[theme]
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
        st = self.d.styles["Normal"]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.15
        for name, size, color in (("Heading 1", 20, self.t["pri"]), ("Heading 2", 15, self.t["pri"]), ("Heading 3", 12.5, self.t["acc"] if self.t["acc"] != "C9A227" else "8A6A00")):
            h = self.d.styles[name]
            h.font.name = self.t["head"]
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = rgb(color)
            h.element.rPr.rFonts.set(qn("w:ascii"), self.t["head"])
            h.element.rPr.rFonts.set(qn("w:hAnsi"), self.t["head"])
            h.paragraph_format.space_before = Pt(16 if name == "Heading 1" else 12)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        self._footer(header_text)

    def _footer(self, text):
        sec = self.d.sections[0]
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text + ("  ·  Página " if text else "Página "))
        r.font.size = Pt(9)
        r.font.color.rgb = rgb("777777")
        r2 = p.add_run()
        r2.font.size = Pt(9)
        r2.font.color.rgb = rgb("777777")
        for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if typ:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), typ)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
            r2._r.append(e)

    # -------- bloques --------
    def cover(self, title, subtitle, meta=None, logo=None, logo_w=9.0):
        if logo:
            p = self.d.add_paragraph()
            p.paragraph_format.space_before = Pt(30)
            p.add_run().add_picture(logo, width=Cm(logo_w))
        p = self.d.add_paragraph()
        p.paragraph_format.space_before = Pt(90 if logo else 120)
        r = p.add_run(title)
        r.font.size = Pt(34)
        r.font.bold = True
        r.font.name = self.t["head"]
        r.font.color.rgb = rgb(self.t["pri"])
        p = self.d.add_paragraph()
        r = p.add_run(subtitle)
        r.font.size = Pt(16)
        r.font.color.rgb = rgb(self.t["acc"] if self.t["acc"] != "C9A227" else "8A6A00")
        r.font.name = self.t["head"]
        if meta:
            p = self.d.add_paragraph()
            p.paragraph_format.space_before = Pt(60)
            for i, m in enumerate(meta):
                r = p.add_run(m + ("\n" if i < len(meta) - 1 else ""))
                r.font.size = Pt(11)
                r.font.color.rgb = rgb("555555")
        self.d.add_page_break()

    def h1(self, t):
        return self.d.add_heading(t, 1)

    def h2(self, t):
        return self.d.add_heading(t, 2)

    def h3(self, t):
        return self.d.add_heading(t, 3)

    def p(self, text, bold=False, italic=False, size=None, color=None, align=None, after=None):
        para = self.d.add_paragraph()
        r = para.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        if align == "c":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if after is not None:
            para.paragraph_format.space_after = Pt(after)
        return para

    def rich(self, parts):
        """parts: lista de (texto, bold)."""
        para = self.d.add_paragraph()
        for t, b in parts:
            r = para.add_run(t)
            r.bold = b
        return para

    def bullets(self, items, style="List Bullet"):
        for it in items:
            para = self.d.add_paragraph(style=style)
            if isinstance(it, tuple):
                r = para.add_run(it[0])
                r.bold = True
                para.add_run(it[1])
            else:
                para.add_run(it)
            para.paragraph_format.space_after = Pt(3)

    def numbered(self, items):
        self.bullets(items, style="List Number")

    def table(self, rows, widths=None, size=10, header=True, zebra=True, bold_first_col=False, align_right_from=None):
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        borders(t)
        set_cell_margins(t)
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                c = t.cell(ri, ci)
                if widths:
                    c.width = Cm(widths[ci])
                c.text = ""
                para = c.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                r = para.add_run(str(val))
                r.font.size = Pt(size)
                if align_right_from is not None and ci >= align_right_from and ri > 0:
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if header and ri == 0:
                    r.bold = True
                    r.font.color.rgb = rgb("FFFFFF")
                    shade(c, self.t["pri"])
                else:
                    if bold_first_col and ci == 0:
                        r.bold = True
                    if zebra and ri % 2 == 0:
                        shade(c, self.t["soft"])
        if header:
            trPr = t.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            trPr.append(th)
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            trPr.append(cs)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def callout(self, title, text, kind="info"):
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        fill = {"info": self.t["soft"], "warn": "FFF4D6", "alert": "FDECEE"}[kind]
        c = t.cell(0, 0)
        shade(c, fill)
        borders(t, "BBBBBB")
        set_cell_margins(t, 100, 100, 160, 160)
        c.text = ""
        para = c.paragraphs[0]
        r = para.add_run(title + ". ")
        r.bold = True
        r.font.size = Pt(10.5)
        r2 = para.add_run(text)
        r2.font.size = Pt(10.5)
        para.paragraph_format.space_after = Pt(0)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)

    def image(self, path, width_cm=16.0, caption=None):
        para = self.d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(path, width=Cm(width_cm))
        para.paragraph_format.keep_with_next = bool(caption)
        if caption:
            c = self.d.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = c.add_run(caption)
            r.italic = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = rgb("666666")

    def page_break(self):
        self.d.add_page_break()

    def save(self, path):
        self.d.save(path)
        return path

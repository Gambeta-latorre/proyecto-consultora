import sys
sys.path.insert(0, "_build")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datos import SUP, INVERSION, calcular

G = "1F4D3A"
GOLD = "C9A227"
hdr_fill = PatternFill("solid", fgColor=G)
tot_fill = PatternFill("solid", fgColor="E8F1EB")
in_font = Font(color="0000FF")
thin = Side(style="thin", color="C8D3CB")
box = Border(top=thin, bottom=thin, left=thin, right=thin)
USD = '"USD" #,##0;[Red]-"USD" #,##0'
USD2 = '"USD" #,##0.00'
PCT = "0.0%"


def head(ws, row, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=16, color=G)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="666666")


def widths(ws, ws_w):
    for i, w in enumerate(ws_w, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()
# ---------------- Supuestos ----------------
ws = wb.active
ws.title = "Supuestos"
title(ws, "Proyecto ASHAB · Supuestos del modelo", "Las celdas en azul se pueden modificar: todo el libro se recalcula solo. Valores en USD salvo indicación.")
head(ws, 3, ["Supuesto", "Valor", "Nota"])
ref = {}
for i, (k, lab, val, nota) in enumerate(SUP):
    r = 4 + i
    ws.cell(r, 1, lab).border = box
    c = ws.cell(r, 2, val)
    c.font = in_font
    c.border = box
    ws.cell(r, 3, nota).border = box
    if isinstance(val, float) and val < 1:
        c.number_format = "0.0%" if k not in ("c_pack",) else USD2
    elif k in ("c_yerba", "c_pack", "p_cif"):
        c.number_format = USD2
    elif k in ("kg", "n1", "n2", "n3", "tc", "aseg_k"):
        c.number_format = "#,##0.00" if k == "aseg_k" else "#,##0"
    else:
        c.number_format = USD
    ref[k] = f"Supuestos!$B${r}"
ws["B13"].number_format = PCT  # derechos
widths(ws, [62, 16, 70])
ws.freeze_panes = "A4"

# ---------------- Costo por contenedor ----------------
ws = wb.create_sheet("Costo por contenedor")
title(ws, "Costo y margen de un contenedor 40' (24 toneladas)", "Del precio de planta al cobro en destino")
head(ws, 3, ["Concepto", "USD", "USD por kg"])
rows = [
    ("Yerba empacada en planta", f"={ref['kg']}*{ref['c_yerba']}"),
    ("Packaging de marca bilingüe", f"={ref['kg']}*{ref['c_pack']}"),
    ("Flete interno a Buenos Aires", f"={ref['flete_int']}"),
    ("Gastos de exportación", f"={ref['g_exp']}"),
    ("Derechos de exportación", "=SUM(B4:B7)*" + ref["derechos"]),
    ("COSTO FOB", "=SUM(B4:B8)"),
    ("Flete marítimo", f"={ref['flete_mar']}"),
    ("Seguro de carga", f"=(B9+B10)*{ref['aseg_k']}*{ref['seguro']}"),
    ("COSTO CIF", "=B9+B10+B11"),
    ("Venta CIF al importador", f"={ref['kg']}*{ref['p_cif']}"),
    ("Carta de crédito", f"=B13*{ref['lc']}"),
    ("Gastos bancarios y cobranza", f"={ref['banc']}"),
    ("Honorario de éxito de Rumbo Global", f"=B13*{ref['com']}"),
    ("COSTO TOTAL", "=B12+B14+B15+B16"),
    ("MARGEN NETO POR CONTENEDOR", "=B13-B17"),
    ("Margen sobre ventas", "=B18/B13"),
    ("Precio FOB equivalente (USD/kg)", f"=(B13-B10-B11)/{ref['kg']}"),
]
for i, (lab, f) in enumerate(rows):
    r = 4 + i
    ws.cell(r, 1, lab).border = box
    c = ws.cell(r, 2, f)
    c.border = box
    c.number_format = USD
    if lab.startswith("Margen sobre"):
        c.number_format = PCT
    elif lab.startswith("Precio FOB"):
        c.number_format = USD2
    else:
        k = ws.cell(r, 3, f"=B{r}/{ref['kg']}")
        k.number_format = USD2
        k.border = box
    if lab.isupper() or lab.startswith("COSTO") or lab.startswith("MARGEN") or lab.startswith("Venta"):
        for col in (1, 2, 3):
            ws.cell(r, col).font = Font(bold=True)
            ws.cell(r, col).fill = tot_fill
widths(ws, [46, 18, 16])

# ---------------- Inversión inicial ----------------
ws = wb.create_sheet("Inversión inicial")
title(ws, "Inversión inicial del proyecto", "Se paga una sola vez durante el primer año")
head(ws, 3, ["Rubro", "USD", "Detalle"])
for i, (lab, usd, det) in enumerate(INVERSION):
    r = 4 + i
    ws.cell(r, 1, lab).border = box
    c = ws.cell(r, 2, usd)
    c.number_format = USD
    c.font = in_font
    c.border = box
    ws.cell(r, 3, det).border = box
n = len(INVERSION)
r0 = 4 + n
ws.cell(r0, 1, "Subtotal")
ws.cell(r0, 2, f"=SUM(B4:B{r0-1})")
ws.cell(r0 + 1, 1, "Imprevistos")
ws.cell(r0 + 1, 2, f"=B{r0}*{ref['imprev']}")
ws.cell(r0 + 2, 1, "INVERSIÓN INICIAL TOTAL (USD)")
ws.cell(r0 + 2, 2, f"=B{r0}+B{r0+1}")
ws.cell(r0 + 3, 1, "Inversión inicial total (ARS)")
ws.cell(r0 + 3, 2, f"=B{r0+2}*{ref['tc']}")
ws.cell(r0 + 4, 1, "Capital de trabajo recomendado (2 contenedores, con seña del 50 %)")
ws.cell(r0 + 4, 2, f"=('Costo por contenedor'!B17-'Costo por contenedor'!B13*{ref['sena']})*2")
ws.cell(r0 + 5, 1, "NECESIDAD MÁXIMA DE FINANCIAMIENTO (USD)")
ws.cell(r0 + 5, 2, f"=B{r0+2}+B{r0+4}")
for rr in range(r0, r0 + 6):
    ws.cell(rr, 2).number_format = USD if rr != r0 + 3 else '"ARS" #,##0'
    for col in (1, 2):
        ws.cell(rr, col).border = box
        if rr in (r0 + 2, r0 + 5):
            ws.cell(rr, col).font = Font(bold=True)
            ws.cell(rr, col).fill = tot_fill
widths(ws, [62, 18, 60])
INV_TOTAL = f"'Inversión inicial'!$B${r0+2}"

# ---------------- Proyección ----------------
ws = wb.create_sheet("Proyección 3 años")
title(ws, "Proyección a 3 años", "No incluye impuesto a las ganancias ni costo financiero. El IVA de exportación está exento (se recupera el crédito fiscal).")
head(ws, 3, ["Concepto", "Año 1", "Año 2", "Año 3"])
lines = [
    ("Contenedores de 40'", [f"={ref['n1']}", f"={ref['n2']}", f"={ref['n3']}"], "#,##0"),
    ("Ingresos por ventas", ["=B4*'Costo por contenedor'!$B$13", "=C4*'Costo por contenedor'!$B$13", "=D4*'Costo por contenedor'!$B$13"], USD),
    ("Costos variables (producto, logística, comisiones)", ["=B4*'Costo por contenedor'!$B$17", "=C4*'Costo por contenedor'!$B$17", "=D4*'Costo por contenedor'!$B$17"], USD),
    ("Margen de contribución", ["=B5-B6", "=C5-C6", "=D5-D6"], USD),
    ("Costos fijos (web, agente, feria, seguimiento)", [0, f"={ref['f_web']}+{ref['f_agente']}+{ref['f_feria']}+{ref['f_seg']}", f"={ref['f_web']}+{ref['f_agente']}+{ref['f_feria']}+{ref['f_seg']}"], USD),
    ("Inversión inicial", [f"={INV_TOTAL}", 0, 0], USD),
    ("FLUJO NETO DEL AÑO", ["=B7-B8-B9", "=C7-C8-C9", "=D7-D8-D9"], USD),
    ("FLUJO ACUMULADO", ["=B10", "=B11+C10", "=C11+D10"], USD),
]
for i, (lab, fs, nf) in enumerate(lines):
    r = 4 + i
    ws.cell(r, 1, lab).border = box
    for j, f in enumerate(fs):
        c = ws.cell(r, 2 + j, f)
        c.number_format = nf
        c.border = box
    if lab.isupper() or lab.startswith("FLUJO"):
        for col in range(1, 5):
            ws.cell(r, col).font = Font(bold=True)
            ws.cell(r, col).fill = tot_fill
ws["A13"] = "Recupero de la inversión (meses aprox.)"
ws["B13"] = "=IF(B11>=0,12*B9/B7,12+(-B11)/(C10/12))"
ws["B13"].number_format = "0.0"
ws["A13"].font = Font(bold=True)
widths(ws, [52, 20, 20, 20])

# ---------------- Escenarios ----------------
ws = wb.create_sheet("Escenarios")
title(ws, "Escenarios a 3 años", "Se modifican el precio de venta y el volumen respecto del caso base")
head(ws, 3, ["Escenario", "Variación de precio", "Variación de volumen", "Margen por contenedor", "Flujo acumulado a 3 años"])
esc = [("Pesimista", -0.10, -0.30), ("Base", 0.0, 0.0), ("Optimista", 0.05, 0.20)]
for i, (nom, dp, dv) in enumerate(esc):
    r = 4 + i
    ws.cell(r, 1, nom)
    c = ws.cell(r, 2, dp)
    c.number_format = PCT
    c.font = in_font
    c = ws.cell(r, 3, dv)
    c.number_format = PCT
    c.font = in_font
    ws.cell(r, 4, f"={ref['kg']}*{ref['p_cif']}*(1+B{r})*(1-{ref['lc']}-{ref['com']})-('Costo por contenedor'!$B$12+'Costo por contenedor'!$B$15)").number_format = USD
    ws.cell(r, 5, f"=D{r}*({ref['n1']}+{ref['n2']}+{ref['n3']})*(1+C{r})-{INV_TOTAL}-2*({ref['f_web']}+{ref['f_agente']}+{ref['f_feria']}+{ref['f_seg']})").number_format = USD
    for col in range(1, 6):
        ws.cell(r, col).border = box
widths(ws, [18, 20, 22, 26, 28])

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb.save("06_Documentos/Costos_del_Proyecto_ASHAB.xlsx")
print("xlsx ok")
